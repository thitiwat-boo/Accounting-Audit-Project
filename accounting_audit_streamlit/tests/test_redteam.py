import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from src.mapping import suggest_mapping
from src.validators import tb_quality, gl_quality, duplicate_gl_rows
from src.financial_statements import build_statement
from src.exporters import export_workbook, REQUIRED_SHEETS

def sample_tb():
    return pd.DataFrame([
        {'Account Code':'110101','Account Name':'เงินสดและเงินฝากธนาคาร','Debit':100000,'Credit':0,'_Source_Row':2},
        {'Account Code':'210101','Account Name':'เจ้าหนี้การค้า','Debit':0,'Credit':20000,'_Source_Row':3},
        {'Account Code':'310101','Account Name':'ทุนจดทะเบียน','Debit':0,'Credit':50000,'_Source_Row':4},
        {'Account Code':'410101','Account Name':'รายได้จากการให้บริการ','Debit':0,'Credit':80000,'_Source_Row':5},
        {'Account Code':'510101','Account Name':'ต้นทุนบริการ','Debit':50000,'Credit':0,'_Source_Row':6},
    ])
def test_tc01_mapping_and_statement_gate():
    tb=sample_tb(); m=suggest_mapping(tb); assert len(m)==len(tb); assert m['Reviewer Status'].eq('Pending').all()
    m['Reviewer Status']='Approved'; r=build_statement(tb,m); assert not r['blocked']; assert len(r['income_statement'])>0

def test_tc02_tb_unbalanced_critical():
    tb=sample_tb(); tb.loc[0,'Debit']=90000; q=tb_quality(tb); assert (q.loc[q.Check.eq('TB Debit = Credit'),'Status'].iloc[0]=='Critical')

def test_tc03_paired_journal_not_duplicate():
    gl=pd.DataFrame([{'Date':'2025-12-31','Document No':'JV1','Account Code':'1101','Debit':100,'Credit':0,'Description':'expense','Reference':'X'}, {'Date':'2025-12-31','Document No':'JV1','Account Code':'5101','Debit':0,'Credit':100,'Description':'expense','Reference':'X'}])
    assert duplicate_gl_rows(gl).empty

def test_tc04_suspense_not_safe_approved():
    tb=pd.DataFrame([{'Account Code':'9999','Account Name':'บัญชีพัก Suspense','Debit':100,'Credit':100}]); m=suggest_mapping(tb); assert m.iloc[0]['Reviewer Status']=='Pending'; assert m.iloc[0]['Confidence']<.7

def test_tc05_unmapped_blocks():
    tb=sample_tb(); tb.loc[len(tb)]={'Account Code':'9999','Account Name':'Other unknown','Debit':1,'Credit':0,'_Source_Row':7}; m=suggest_mapping(tb); r=build_statement(tb,m); assert r['blocked']; assert len(r['unmapped_or_pending'])>0

def test_tc06_prior_year_difference_is_not_fabricated():
    tb=sample_tb(); m=suggest_mapping(tb); r=build_statement(tb,m); assert 'FS_Changes_Equity' in REQUIRED_SHEETS and r['net_income'] is None

def test_excel_has_required_sheets_and_formulas():
    tb=sample_tb(); m=suggest_mapping(tb); m['Reviewer Status']='Approved'
    tables={'TB_Original':tb,'TB_Normalized':tb,'COA_Mapping':m,'FS_Balance_Sheet':pd.DataFrame(),'FS_Income_Statement':pd.DataFrame()}
    blob=export_workbook(tables,{'Project ID':'TEST','TB SHA-256':'abc'})
    wb=load_workbook(BytesIO(blob),data_only=False)
    assert set(REQUIRED_SHEETS).issubset(set(wb.sheetnames))
    assert any(isinstance(c.value,str) and c.value.startswith('=') for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    assert not any('#REF!' in c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value,str))
